from pathlib import Path
from typing import Any


def normalize_header(text: str) -> str:
    return text.strip().lower().replace(" ", "_")


def guess_activity_name(file_path: Path) -> str:
    return file_path.stem.replace("_", " ").title()


def build_record(values: dict[str, Any], source_file: Path) -> dict[str, Any]:
    return {
        "source_file": source_file.name,
        "activity_name": guess_activity_name(source_file),
        "student_id": values.get("student_id") or values.get("id") or values.get("mssv"),
        "student_name": values.get("student_name") or values.get("name"),
        "class": values.get("class") or values.get("grade"),
        "score": values.get("score"),
    }


def parse_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to parse XLSX documents") from exc

    workbook = openpyxl.load_workbook(path, data_only=True)
    records: list[dict[str, Any]] = []

    for sheet in workbook.worksheets:
        headers = []
        for row in sheet.iter_rows(values_only=True):
            if not any(row):
                continue
            row_values = [str(cell).strip() if cell is not None else "" for cell in row]
            normalized = [normalize_header(value) for value in row_values]
            if any(key in normalized for key in ["name", "student_name", "id", "mssv", "score"]):
                headers = normalized
                continue
            if not headers:
                continue

            values = {}
            for header, cell in zip(headers, row_values):
                if header:
                    values[header] = cell
            if values:
                records.append(build_record(values, path))

    return records


def parse_docx(path: Path) -> list[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to parse DOCX documents") from exc

    document = Document(path)
    records: list[dict[str, Any]] = []

    for table in document.tables:
        headers = []
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            normalized = [normalize_header(value) for value in cells]
            if any(key in normalized for key in ["name", "student_name", "id", "mssv", "score"]):
                headers = normalized
                continue
            if not headers:
                continue

            values = {}
            for header, cell_text in zip(headers, cells):
                if header:
                    values[header] = cell_text
            if values:
                records.append(build_record(values, path))

    return records


def parse_documents(file_paths: list[str]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []

    for file_path in file_paths:
        path = Path(file_path)
        if not path.exists():
            continue
        if path.suffix.lower() == ".xlsx":
            records.extend(parse_xlsx(path))
        elif path.suffix.lower() == ".docx":
            records.extend(parse_docx(path))

    return records
